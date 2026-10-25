"""Lê a planilha sem alterá-la e prepara casos para dois motores independentes."""
import csv
import hashlib
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
from monitor.calculos import calcular_linha, consolidar
from monitor.indicadores import extrair

BASE = Path(__file__).resolve().parents[1]


def main():
    fonte = BASE / "referencias/estimativa voto 2T.xlsx"
    pasta = BASE / "evidencias/auditoria"
    pasta.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(fonte, data_only=False)
    cache = openpyxl.load_workbook(fonte, data_only=True)
    ws = wb["Planilha1"]
    formulas = {c.coordinate: c.value for row in ws for c in row if c.data_type == "f"}
    esperado = {f"{col}{r}" for r in range(2, 30) for col in "IJKL"} | {"H30", "J30", "K30", "L30", "M30", "N30"}
    if set(formulas) != esperado:
        raise ValueError("Fórmulas da referência mudaram; revisar a tradução")
    for r in range(2, 30):
        for col, expr in {"I": f"=F{r}/SUM(F{r}:G{r})", "J": f"=H{r}*I{r}*D{r}",
                          "K": f"=H{r}*I{r}*E{r}", "L": f"=H{r}*I{r}"}.items():
            assert formulas[f"{col}{r}"] == expr
    assert {a: formulas[a] for a in ("H30", "J30", "K30", "L30", "M30", "N30")} == {
        "H30": "=SUM(H2:H29)", "J30": "=SUM(J2:J29)", "K30": "=SUM(K2:K29)",
        "L30": "=SUM(L2:L29)", "M30": "=J30/L30", "N30": "=K30/L30"}
    inventario = {"fonte": str(fonte.relative_to(BASE)), "sha256": hashlib.sha256(fonte.read_bytes()).hexdigest(),
        "sheet": ws.title, "formulas": formulas,
        "ufs": [ws[f"C{r}"].value for r in range(2, 30)],
        "entradas_vazias": [f"{c}{r}" for r in range(2, 30) for c in "DEFGH" if ws[f"{c}{r}"].value is None],
        "erros_cache": {c.coordinate: c.value for row in cache[ws.title] for c in row
                        if c.data_type == "e" or isinstance(c.value, str) and c.value.startswith("#")}}
    (pasta / "inventario_formulas.json").write_text(json.dumps(inventario, ensure_ascii=False, indent=2), encoding="utf-8")
    with (pasta / "inventario_formulas.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";"); w.writerow(["aba", "celula", "formula_original", "valor_cache"])
        for a, expr in formulas.items(): w.writerow([ws.title, a, expr, cache[ws.title][a].value])
    casos = [
        ("ficticio_uniforme", [["0.4", "0.6", 80, 20, 100] for _ in range(28)]),
        ("ficticio_heterogeneo_parcial", [[str((i+1)/100), str((80-i)/100), 123+i*37, 20+i, 1000+i*100] for i in range(28)]),
        ("ficticio_comparecimento_total", [["0.33333333333", "0.66666666667", 100001+i, 0, 100001+i] for i in range(28)]),
        ("ficticio_candidato_zero", [["0", "1", 1+i, 100, 101+i] for i in range(28)]),
    ]
    real = extrair(json.loads((BASE / "evidencias/tse_acre_6257.json").read_bytes()), "6257", "1")
    entrada = [str(real["candidatos"][n]["fracao_calculo"]) for n in ("lula", "flavio")] + [real[k] for k in ("comparecimento", "abstencao", "eleitorado_apto")]
    casos.append(("acre_1T_replicado_28_vezes_APENAS_TESTE", [entrada[:] for _ in range(28)]))
    saida = []
    for nome, entradas in casos:
        resultados = {}
        for r, linha in enumerate(entradas, 2):
            calculado = calcular_linha(*linha)
            resultados.update({f"{col}{r}": str(calculado[col]) for col in "IJKL"})
        consolidado = consolidar(entradas)
        resultados.update({cel: str(v) for cel, v in consolidado.items() if cel != "status"})
        saida.append({"nome": nome, "entradas": entradas, "python": resultados})
    (pasta / "casos_equivalencia.json").write_text(json.dumps(saida, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"formulas": len(formulas), "entradas_vazias": len(inventario["entradas_vazias"]),
                      "erros_cache": len(inventario["erros_cache"]), "casos": len(saida), "sha256": inventario["sha256"]}))


if __name__ == "__main__":
    main()
